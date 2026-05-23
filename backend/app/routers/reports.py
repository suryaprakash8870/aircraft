import io
import uuid
from datetime import date, datetime, timezone
from typing import Optional

from fastapi import APIRouter, Depends, HTTPException, status
from fastapi.responses import StreamingResponse
from sqlalchemy import func, select
from sqlalchemy.orm import selectinload
from sqlalchemy.ext.asyncio import AsyncSession

from app.database import get_db
from app.models.aircraft import Aircraft
from app.models.aircraft_filling import AircraftFilling
from app.models.airport import Airport
from app.models.fuel_agent import FuelAgent
from app.models.fuel_purchase import FuelPurchase
from app.models.fuel_stock import FuelStock, FuelStockLog
from app.models.user import User
from app.utils.auth import get_current_active_user

router = APIRouter(prefix="/api/reports", tags=["Reports"])


@router.get("/fuel-purchases")
async def report_fuel_purchases(
    start_date: Optional[date] = None,
    end_date: Optional[date] = None,
    airport_id: Optional[uuid.UUID] = None,
    agent_id: Optional[uuid.UUID] = None,
    format: str = "json",
    current_user: User = Depends(get_current_active_user),
    db: AsyncSession = Depends(get_db),
):
    query = select(FuelPurchase).options(
        selectinload(FuelPurchase.fuel_agent), selectinload(FuelPurchase.airport)
    )
    if start_date:
        query = query.where(FuelPurchase.purchase_date >= start_date)
    if end_date:
        query = query.where(FuelPurchase.purchase_date <= end_date)
    if airport_id:
        query = query.where(FuelPurchase.airport_id == airport_id)
    if agent_id:
        query = query.where(FuelPurchase.fuel_agent_id == agent_id)
    query = query.order_by(FuelPurchase.purchase_date.desc())

    result = await db.execute(query)
    purchases = result.scalars().all()

    data = [
        {
            "purchase_id": p.purchase_id,
            "purchase_date": str(p.purchase_date),
            "fuel_type": p.fuel_type,
            "quantity_purchased": p.quantity_purchased,
            "purchase_rate": p.purchase_rate,
            "total_amount": p.total_amount,
            "payment_status": p.payment_status,
            "invoice_number": p.invoice_number,
            "agent_name": p.fuel_agent.agent_name if p.fuel_agent else None,
            "airport_name": p.airport.airport_name if p.airport else None,
            "airport_code": p.airport.airport_code if p.airport else None,
        }
        for p in purchases
    ]

    if format == "xlsx":
        return _generate_xlsx_response(
            data,
            headers=[
                "purchase_id", "purchase_date", "fuel_type", "quantity_purchased",
                "purchase_rate", "total_amount", "payment_status", "invoice_number",
                "agent_name", "airport_name", "airport_code",
            ],
            filename="fuel_purchases_report.xlsx",
        )
    return {"data": data, "total": len(data)}


@router.get("/fuel-consumption")
async def report_fuel_consumption(
    start_date: Optional[date] = None,
    end_date: Optional[date] = None,
    aircraft_id: Optional[uuid.UUID] = None,
    airport_id: Optional[uuid.UUID] = None,
    format: str = "json",
    current_user: User = Depends(get_current_active_user),
    db: AsyncSession = Depends(get_db),
):
    query = select(AircraftFilling).options(
        selectinload(AircraftFilling.aircraft), selectinload(AircraftFilling.airport)
    )
    if start_date:
        query = query.where(AircraftFilling.filling_datetime >= start_date)
    if end_date:
        query = query.where(AircraftFilling.filling_datetime <= end_date)
    if aircraft_id:
        query = query.where(AircraftFilling.aircraft_id == aircraft_id)
    if airport_id:
        query = query.where(AircraftFilling.airport_id == airport_id)
    query = query.order_by(AircraftFilling.filling_datetime.desc())

    result = await db.execute(query)
    fillings = result.scalars().all()

    data = [
        {
            "filling_id": f.filling_id,
            "filling_datetime": str(f.filling_datetime),
            "aircraft_number": f.aircraft.aircraft_number if f.aircraft else None,
            "airline_name": f.aircraft.airline_name if f.aircraft else None,
            "flight_number": f.flight_number,
            "airport_code": f.airport.airport_code if f.airport else None,
            "airport_name": f.airport.airport_name if f.airport else None,
            "quantity_filled": f.quantity_filled,
            "fuel_rate": f.fuel_rate,
            "total_cost": f.total_cost,
        }
        for f in fillings
    ]

    if format == "xlsx":
        return _generate_xlsx_response(
            data,
            headers=[
                "filling_id", "filling_datetime", "aircraft_number", "airline_name",
                "flight_number", "airport_code", "airport_name", "quantity_filled",
                "fuel_rate", "total_cost",
            ],
            filename="fuel_consumption_report.xlsx",
        )
    return {"data": data, "total": len(data)}


@router.get("/airport-stock")
async def report_airport_stock(
    airport_id: Optional[uuid.UUID] = None,
    current_user: User = Depends(get_current_active_user),
    db: AsyncSession = Depends(get_db),
):
    query = select(FuelStock).options(selectinload(FuelStock.airport))
    if airport_id:
        query = query.where(FuelStock.airport_id == airport_id)
    query = query.order_by(FuelStock.last_updated.desc())

    result = await db.execute(query)
    stocks = result.scalars().all()

    data = [
        {
            "airport_name": s.airport.airport_name if s.airport else None,
            "airport_code": s.airport.airport_code if s.airport else None,
            "fuel_type": s.fuel_type,
            "current_stock": s.current_stock,
            "fuel_storage_capacity": s.airport.fuel_storage_capacity if s.airport else None,
            "last_updated": str(s.last_updated),
        }
        for s in stocks
    ]
    return {"data": data, "total": len(data)}


@router.get("/aircraft-history")
async def report_aircraft_history(
    aircraft_id: Optional[uuid.UUID] = None,
    start_date: Optional[date] = None,
    end_date: Optional[date] = None,
    format: str = "json",
    current_user: User = Depends(get_current_active_user),
    db: AsyncSession = Depends(get_db),
):
    if not aircraft_id:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="aircraft_id is required",
        )

    aircraft_result = await db.execute(select(Aircraft).where(Aircraft.id == aircraft_id))
    aircraft = aircraft_result.scalar_one_or_none()
    if not aircraft:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Aircraft not found")

    query = select(AircraftFilling).options(selectinload(AircraftFilling.airport)).where(
        AircraftFilling.aircraft_id == aircraft_id
    )
    if start_date:
        query = query.where(AircraftFilling.filling_datetime >= start_date)
    if end_date:
        query = query.where(AircraftFilling.filling_datetime <= end_date)
    query = query.order_by(AircraftFilling.filling_datetime.desc())

    result = await db.execute(query)
    fillings = result.scalars().all()

    total_qty = sum(f.quantity_filled for f in fillings)
    total_cost = sum(f.total_cost for f in fillings)

    data = {
        "aircraft": {
            "id": str(aircraft.id),
            "aircraft_number": aircraft.aircraft_number,
            "aircraft_model": aircraft.aircraft_model,
            "airline_name": aircraft.airline_name,
        },
        "summary": {
            "total_fillings": len(fillings),
            "total_quantity_filled": total_qty,
            "total_cost": total_cost,
        },
        "fillings": [
            {
                "filling_id": f.filling_id,
                "filling_datetime": str(f.filling_datetime),
                "airport_code": f.airport.airport_code if f.airport else None,
                "airport_name": f.airport.airport_name if f.airport else None,
                "flight_number": f.flight_number,
                "quantity_filled": f.quantity_filled,
                "fuel_rate": f.fuel_rate,
                "total_cost": f.total_cost,
            }
            for f in fillings
        ],
    }
    return data


@router.get("/vendor")
async def report_vendor(
    agent_id: Optional[uuid.UUID] = None,
    start_date: Optional[date] = None,
    end_date: Optional[date] = None,
    format: str = "json",
    current_user: User = Depends(get_current_active_user),
    db: AsyncSession = Depends(get_db),
):
    if not agent_id:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="agent_id is required",
        )

    agent_result = await db.execute(select(FuelAgent).where(FuelAgent.id == agent_id))
    agent = agent_result.scalar_one_or_none()
    if not agent:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Fuel agent not found")

    query = select(FuelPurchase).options(selectinload(FuelPurchase.airport)).where(
        FuelPurchase.fuel_agent_id == agent_id
    )
    if start_date:
        query = query.where(FuelPurchase.purchase_date >= start_date)
    if end_date:
        query = query.where(FuelPurchase.purchase_date <= end_date)
    query = query.order_by(FuelPurchase.purchase_date.desc())

    result = await db.execute(query)
    purchases = result.scalars().all()

    total_qty = sum(p.quantity_purchased for p in purchases)
    total_amount = sum(p.total_amount for p in purchases)

    data = {
        "agent": {
            "id": str(agent.id),
            "agent_name": agent.agent_name,
            "company_name": agent.company_name,
            "gst_number": agent.gst_number,
        },
        "summary": {
            "total_purchases": len(purchases),
            "total_quantity": total_qty,
            "total_amount": total_amount,
        },
        "purchases": [
            {
                "purchase_id": p.purchase_id,
                "purchase_date": str(p.purchase_date),
                "airport_code": p.airport.airport_code if p.airport else None,
                "fuel_type": p.fuel_type,
                "quantity_purchased": p.quantity_purchased,
                "purchase_rate": p.purchase_rate,
                "total_amount": p.total_amount,
                "payment_status": p.payment_status,
                "invoice_number": p.invoice_number,
            }
            for p in purchases
        ],
    }
    return data


@router.get("/export")
async def export_report(
    type: str = "purchases",
    format: str = "xlsx",
    start_date: Optional[date] = None,
    end_date: Optional[date] = None,
    airport_id: Optional[uuid.UUID] = None,
    agent_id: Optional[uuid.UUID] = None,
    current_user: User = Depends(get_current_active_user),
    db: AsyncSession = Depends(get_db),
):
    if type == "purchases":
        query = select(FuelPurchase).options(
            selectinload(FuelPurchase.fuel_agent), selectinload(FuelPurchase.airport)
        )
        if start_date:
            query = query.where(FuelPurchase.purchase_date >= start_date)
        if end_date:
            query = query.where(FuelPurchase.purchase_date <= end_date)
        if airport_id:
            query = query.where(FuelPurchase.airport_id == airport_id)
        if agent_id:
            query = query.where(FuelPurchase.fuel_agent_id == agent_id)
        query = query.order_by(FuelPurchase.purchase_date.desc())
        result = await db.execute(query)
        records = result.scalars().all()

        data = [
            {
                "Purchase ID": p.purchase_id,
                "Date": str(p.purchase_date),
                "Fuel Type": p.fuel_type,
                "Quantity (L)": p.quantity_purchased,
                "Rate (per L)": p.purchase_rate,
                "Total Amount": p.total_amount,
                "Payment Status": p.payment_status,
                "Invoice Number": p.invoice_number or "",
                "Agent": p.fuel_agent.agent_name if p.fuel_agent else "",
                "Airport": p.airport.airport_name if p.airport else "",
                "Airport Code": p.airport.airport_code if p.airport else "",
            }
            for p in records
        ]
        filename = f"fuel_purchases_{datetime.now().strftime('%Y%m%d')}.xlsx"

    elif type == "consumption":
        query = select(AircraftFilling).options(
            selectinload(AircraftFilling.aircraft), selectinload(AircraftFilling.airport)
        )
        if start_date:
            query = query.where(AircraftFilling.filling_datetime >= start_date)
        if end_date:
            query = query.where(AircraftFilling.filling_datetime <= end_date)
        if airport_id:
            query = query.where(AircraftFilling.airport_id == airport_id)
        query = query.order_by(AircraftFilling.filling_datetime.desc())
        result = await db.execute(query)
        records = result.scalars().all()

        data = [
            {
                "Filling ID": f.filling_id,
                "Date/Time": str(f.filling_datetime),
                "Aircraft Number": f.aircraft.aircraft_number if f.aircraft else "",
                "Airline": f.aircraft.airline_name if f.aircraft else "",
                "Flight Number": f.flight_number or "",
                "Airport": f.airport.airport_name if f.airport else "",
                "Airport Code": f.airport.airport_code if f.airport else "",
                "Quantity Filled (L)": f.quantity_filled,
                "Fuel Rate": f.fuel_rate,
                "Total Cost": f.total_cost,
            }
            for f in records
        ]
        filename = f"fuel_consumption_{datetime.now().strftime('%Y%m%d')}.xlsx"

    else:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="type must be 'purchases' or 'consumption'",
        )

    if format == "xlsx":
        return _generate_xlsx_response(data, headers=list(data[0].keys()) if data else [], filename=filename)

    raise HTTPException(
        status_code=status.HTTP_400_BAD_REQUEST,
        detail="format must be 'xlsx'",
    )


def _generate_xlsx_response(data: list, headers: list, filename: str) -> StreamingResponse:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Report"

    header_fill = PatternFill(start_color="1A3C6E", end_color="1A3C6E", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_alignment = Alignment(horizontal="center", vertical="center")

    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    if not headers and data:
        headers = list(data[0].keys())

    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border

    alt_fill = PatternFill(start_color="ECF0F1", end_color="ECF0F1", fill_type="solid")
    for row_idx, row_data in enumerate(data, start=2):
        for col_idx, header in enumerate(headers, start=1):
            value = row_data.get(header, "") if isinstance(row_data, dict) else row_data[col_idx - 1]
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.border = thin_border
            cell.alignment = Alignment(vertical="center")
            if row_idx % 2 == 0:
                cell.fill = alt_fill

    for col_idx in range(1, len(headers) + 1):
        max_length = 0
        col_letter = get_column_letter(col_idx)
        for cell in ws[col_letter]:
            try:
                if cell.value and len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except Exception:
                pass
        ws.column_dimensions[col_letter].width = min(max_length + 4, 40)

    ws.row_dimensions[1].height = 25

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )
